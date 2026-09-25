#!/usr/bin/env python3
"""Update datos/indicadores/<id>.json from plantillas/indicador_<id>.xlsx.

Usage (from the repository root):
    python herramientas/actualizar_indicador.py plantillas/indicador_per_beneficiarios.xlsx
    python herramientas/actualizar_indicador.py plantillas/indicador_per_beneficiarios.xlsx --solo-validar

The Excel replaces the indicator's values entirely (the sheet is the full record).
A new indicator is created if its file does not exist yet, and indice.json is updated.
"""
import argparse
import re
from datetime import date, datetime

from openpyxl import load_workbook

from comun import (DIR_INDICADORES, RUTA_FUENTES, RUTA_INDICE, a_bool, catalogo_entidades, catalogo_municipios,
                   escribir_json, leer_json, respaldar, texto, verificar_raiz)

TEMAS = {"migracion", "desplazamiento", "desaparicion", "periodistas"}
CAMPOS_DEF = ["id", "nombre", "tema", "unidad", "fuente", "nota", "tipo"]


def leer_definicion(wb):
    ws = wb["definicion"]
    d = {}
    for fila in ws.iter_rows(values_only=True):
        if fila and fila[0]:
            d[str(fila[0]).strip()] = texto(fila[1]) if len(fila) > 1 else ""
    return {k: d.get(k, "") for k in CAMPOS_DEF}


def periodo_texto(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = texto(v)
    if re.match(r"^\d{4}(-\d{2}(-\d{2})?)?$", s):
        return s
    if re.match(r"^\d{4}\.0$", s):
        return s[:4]
    return None


def leer_valores(wb, ents, muns):
    ws = wb["valores"]
    filas = list(ws.iter_rows(values_only=True))
    enc = [texto(c).lower() for c in filas[0]]
    idx = {c: enc.index(c) for c in ["cve_ent", "cve_mun", "periodo", "valor", "ejemplo", "nota"]}
    valores, errores, vistos = [], [], set()
    for n, fila in enumerate(filas[1:], start=2):
        celdas = {c: fila[i] for c, i in idx.items()}
        if all(texto(celdas[c]) == "" for c in ("cve_ent", "cve_mun", "periodo", "valor")):
            continue
        e = []
        cve = texto(celdas["cve_ent"]).zfill(2)
        if cve not in ents:
            e.append(f"cve_ent '{cve}' no válida")
        cm = texto(celdas["cve_mun"])
        if cm:
            cm = cm.zfill(3)
            if (cve, cm) not in muns:
                e.append(f"cve_mun '{cm}' no está en el catálogo de municipios de la entidad {cve}")
        periodo = periodo_texto(celdas["periodo"])
        if not periodo:
            e.append("periodo no válido (usar AAAA, AAAA-MM o AAAA-MM-DD)")
        try:
            valor = float(celdas["valor"])
            valor = int(valor) if valor.is_integer() else valor
        except (TypeError, ValueError):
            valor = None
            e.append("valor debe ser numérico")
        clave = (cve, cm, periodo)
        if clave in vistos:
            e.append("fila repetida (misma entidad, municipio y periodo)")
        vistos.add(clave)
        if e:
            errores.append(f"Fila {n}: " + "; ".join(e))
            continue
        reg = {"cve_ent": cve, "periodo": periodo, "valor": valor, "ejemplo": a_bool(celdas["ejemplo"])}
        if cm:
            reg["cve_mun"] = cm
        if texto(celdas["nota"]):
            reg["nota"] = texto(celdas["nota"])
        valores.append(reg)
    return valores, errores


def main():
    ap = argparse.ArgumentParser(description="Actualiza un indicador desde su plantilla Excel")
    ap.add_argument("excel")
    ap.add_argument("--solo-validar", action="store_true")
    args = ap.parse_args()
    verificar_raiz()

    wb = load_workbook(args.excel, data_only=True)
    d = leer_definicion(wb)
    errores = []
    if not re.match(r"^[a-z0-9_]+$", d["id"]):
        errores.append("id de la definición vacío o con caracteres no permitidos (usar minúsculas, dígitos y guion bajo)")
    if d["tema"] not in TEMAS:
        errores.append(f"tema '{d['tema']}' no válido")
    fuentes = {f["id"] for f in leer_json(RUTA_FUENTES)}
    if d["fuente"] and d["fuente"] not in fuentes:
        errores.append(f"fuente '{d['fuente']}' no está en datos/fuentes.json")
    ents = {c for c, _ in catalogo_entidades()}
    muns = {(e, m) for e, m, _ in catalogo_municipios()}
    valores, err_val = leer_valores(wb, ents, muns)
    errores += err_val
    if errores:
        print("Errores encontrados, no se escribió nada:")
        print("\n".join("  " + e for e in errores))
        raise SystemExit(1)

    definicion = {k: v for k, v in d.items() if v}
    ruta = DIR_INDICADORES / f"{d['id']}.json"
    print(f"Indicador {d['id']}: {len(valores)} valores válidos.")
    if args.solo_validar:
        print("Modo solo validar, no se escribió nada.")
        return
    resp = respaldar(ruta)
    if resp:
        print(f"Respaldo: {resp}")
    escribir_json(ruta, {"definicion": definicion, "valores": valores})
    indice = leer_json(RUTA_INDICE)
    if ruta.name not in indice["archivos"]:
        indice["archivos"].append(ruta.name)
        escribir_json(RUTA_INDICE, indice)
        print(f"Indicador nuevo añadido a {RUTA_INDICE}")
    print(f"Escrito {ruta}")


if __name__ == "__main__":
    main()
