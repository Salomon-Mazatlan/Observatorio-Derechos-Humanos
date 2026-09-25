#!/usr/bin/env python3
"""Update datos/marco_legal.json from plantillas/marco_legal.xlsx.

Usage (from the repository root):
    python herramientas/actualizar_marco_legal.py plantillas/marco_legal.xlsx
    python herramientas/actualizar_marco_legal.py plantillas/marco_legal.xlsx --solo-validar

The two sheets are the full record: every entity must appear once in 'entidades';
'instrumentos' rows are attached to their entity by cve_ent.
"""
import argparse
import re
from datetime import date, datetime

from openpyxl import load_workbook

from comun import RUTA_MARCO, catalogo_entidades, escribir_json, leer_json, respaldar, texto, verificar_raiz

COLS_ENT = ["cve_ent", "entidad", "categoria", "fecha_corte", "fuente", "nota"]
COLS_INST = ["cve_ent", "nombre", "tipo", "anio", "url", "organo", "nota"]
TIPOS = {"ley", "decreto", "acuerdo", "protocolo", "unidad", "reforma", "iniciativa", "reglamento"}
DESGLOSE_VALORES = {"si": True, "sí": True, "no": False, "parcial": "parcial", "sin dato": None, "": None}


def filas(ws, columnas):
    datos = list(ws.iter_rows(values_only=True))
    enc = [texto(c).lower() for c in datos[0]]
    faltan = [c for c in columnas if c not in enc]
    if faltan:
        raise SystemExit(f"Faltan columnas en la hoja '{ws.title}': {', '.join(faltan)}")
    extra = [c for c in enc if c and c not in columnas]
    out = []
    for n, fila in enumerate(datos[1:], start=2):
        if all(c in (None, "") for c in fila):
            continue
        reg = {c: fila[enc.index(c)] for c in columnas}
        reg["_extra"] = {c: fila[enc.index(c)] for c in extra}
        out.append((n, reg))
    return out


def fecha_texto(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = texto(v)
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else None


def main():
    ap = argparse.ArgumentParser(description="Actualiza datos/marco_legal.json desde Excel")
    ap.add_argument("excel")
    ap.add_argument("--solo-validar", action="store_true")
    args = ap.parse_args()
    verificar_raiz()

    actual = leer_json(RUTA_MARCO)
    categorias = set(actual["categorias"])
    ents = dict(catalogo_entidades())
    wb = load_workbook(args.excel, data_only=True)
    errores = []

    entidades, vistas = {}, set()
    for n, r in filas(wb["entidades"], COLS_ENT):
        e = []
        cve = texto(r["cve_ent"]).zfill(2)
        if cve not in ents:
            e.append(f"cve_ent '{cve}' no válida")
        if cve in vistas:
            e.append(f"entidad {cve} repetida")
        vistas.add(cve)
        cat = texto(r["categoria"])
        if cat not in categorias:
            e.append(f"categoria '{cat}' no válida")
        fc = fecha_texto(r["fecha_corte"]) if texto(r["fecha_corte"]) else actual.get("fecha_corte")
        if texto(r["fecha_corte"]) and not fc:
            e.append("fecha_corte no válida (AAAA-MM-DD)")
        desglose = {}
        for k, v in r["_extra"].items():
            s = texto(v).lower()
            if s not in DESGLOSE_VALORES:
                e.append(f"desglose '{k}' con valor '{s}' no permitido (si, no, parcial, sin dato)")
            desglose[k] = DESGLOSE_VALORES.get(s)
        if e:
            errores.append(f"entidades, fila {n}: " + "; ".join(e))
            continue
        entidades[cve] = {"cve_ent": cve, "entidad": ents[cve], "categoria": cat, "instrumentos": [],
                          "nota": texto(r["nota"]), "fuente": texto(r["fuente"]) or "ishr", "fecha_corte": fc, "desglose": desglose}
    faltantes = sorted(set(ents) - vistas)
    if faltantes:
        errores.append("entidades: faltan " + ", ".join(faltantes))

    for n, r in filas(wb["instrumentos"], COLS_INST):
        e = []
        cve = texto(r["cve_ent"]).zfill(2)
        if cve not in entidades:
            e.append(f"cve_ent '{cve}' sin fila en 'entidades'")
        tipo = texto(r["tipo"]).lower()
        if tipo not in TIPOS:
            e.append(f"tipo '{tipo}' no válido")
        if not texto(r["nombre"]):
            e.append("falta nombre")
        anio = None
        if texto(r["anio"]):
            try:
                anio = int(float(texto(r["anio"])))
                if not 1990 <= anio <= 2100:
                    e.append("anio fuera de rango")
            except ValueError:
                e.append("anio no numérico")
        if e:
            errores.append(f"instrumentos, fila {n}: " + "; ".join(e))
            continue
        entidades[cve]["instrumentos"].append({"nombre": texto(r["nombre"]), "tipo": tipo, "anio": anio,
                                              "url": texto(r["url"]), "organo": texto(r["organo"]), "nota": texto(r["nota"])})
    if errores:
        print("Errores encontrados, no se escribió nada:")
        print("\n".join("  " + x for x in errores))
        raise SystemExit(1)

    total = sum(len(e["instrumentos"]) for e in entidades.values())
    print(f"{len(entidades)} entidades, {total} instrumentos válidos.")
    if args.solo_validar:
        print("Modo solo validar, no se escribió nada.")
        return
    resp = respaldar(RUTA_MARCO)
    if resp:
        print(f"Respaldo: {resp}")
    actual["entidades"] = [entidades[k] for k in sorted(entidades)]
    escribir_json(RUTA_MARCO, actual)
    print(f"Escrito {RUTA_MARCO}")


if __name__ == "__main__":
    main()
