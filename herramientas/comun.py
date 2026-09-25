"""Shared helpers for the capture scripts (paths, catalogs, Excel styling)."""
import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(".")
RUTA_EVENTOS = RAIZ / "datos/eventos.json"
RUTA_FUENTES = RAIZ / "datos/fuentes.json"
RUTA_MARCO = RAIZ / "datos/marco_legal.json"
DIR_INDICADORES = RAIZ / "datos/indicadores"
RUTA_INDICE = DIR_INDICADORES / "indice.json"
RUTA_ESTADOS = RAIZ / "datos/geo/estados.geojson"
RUTA_MUNICIPIOS = RAIZ / "datos/geo/sinaloa_municipios.geojson"
DIR_PLANTILLAS = RAIZ / "plantillas"

ARIAL = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
AMARILLO = PatternFill("solid", fgColor="FFF2CC")
AZUL = PatternFill("solid", fgColor="D9E2F3")


def leer_json(ruta):
    return json.loads(Path(ruta).read_text(encoding="utf-8"))


def escribir_json(ruta, datos):
    Path(ruta).write_text(json.dumps(datos, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def respaldar(ruta):
    ruta = Path(ruta)
    if ruta.exists():
        destino = ruta.with_name(f"{ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}.bak.json")
        shutil.copy(ruta, destino)
        return destino
    return None


def texto(v):
    return "" if v is None else str(v).strip()


def a_bool(v):
    return texto(v).lower() in ("si", "sí", "true", "1", "x", "verdadero")


def verificar_raiz():
    if not RUTA_FUENTES.exists():
        raise SystemExit("Ejecuta el script desde la raíz del repositorio (no encuentro datos/fuentes.json)")


def catalogo_entidades():
    """Returns [(cve_ent, nombre)] from the states GeoJSON."""
    geo = leer_json(RUTA_ESTADOS)
    return sorted((f["properties"]["cve_ent"], f["properties"]["nombre"]) for f in geo["features"])


def catalogo_municipios():
    """Returns [(cve_ent, cve_mun, nombre)] from the Sinaloa GeoJSON."""
    geo = leer_json(RUTA_MUNICIPIOS)
    return sorted((f["properties"]["cve_ent"], f["properties"]["cve_mun"], f["properties"]["nombre"]) for f in geo["features"])


def estilizar_hoja(ws, anchos, filas_editables=0, columnas_ajustables=()):
    """Arial everywhere, bold blue header, yellow input area, column widths."""
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for c in row:
            c.font = ARIAL
            c.alignment = Alignment(vertical="top", wrap_text=c.column in columnas_ajustables)
    for c in ws[1]:
        c.font, c.fill = BOLD, AZUL
    for r in range(2, filas_editables + 2):
        for ci in range(1, len(anchos) + 1):
            ws.cell(r, ci).fill = AMARILLO


def hoja_instrucciones(wb, lineas):
    ws = wb.create_sheet("instrucciones")
    for t in lineas:
        ws.append([t])
    for row in ws.iter_rows():
        for c in row:
            c.font = BOLD if c.row == 1 else ARIAL
    ws.column_dimensions["A"].width = 130
    return ws
