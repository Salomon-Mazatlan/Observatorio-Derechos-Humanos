#!/usr/bin/env python3
"""Build captura_eventos.xlsx from datos/eventos.json and datos/fuentes.json.

Usage (from the repository root):
    python herramientas/crear_plantilla.py [salida.xlsx]
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLS = ["id", "tema", "tipo", "fecha", "lat", "lon", "lugar", "titulo", "descripcion", "fuente", "url", "verificacion", "ejemplo"]
ANCHOS = [9, 16, 22, 12, 10, 11, 30, 48, 70, 16, 40, 15, 9]
TEMAS = ["migracion", "desplazamiento", "desaparicion", "periodistas"]
VERIF = ["oficial", "organización", "campo", "prensa", "sin verificar"]
MAX_FILAS = 1000

INSTRUCCIONES = [
    "Plantilla de captura de eventos del observatorio",
    "",
    "Captura en la hoja 'eventos' (celdas amarillas). Cada fila es un marcador en el mapa.",
    "Las filas con ejemplo = si son plantillas; bórralas o cambia ejemplo = no cuando las sustituyas por hechos reales.",
    "",
    "id: déjalo vacío en registros nuevos, el script lo asigna (ev-001, ev-002...). Para corregir un evento ya publicado conserva su id.",
    "tema: migracion, desplazamiento, desaparicion o periodistas (lista desplegable).",
    "tipo: texto libre pero consistente (desplazamiento masivo, retorno, padrón oficial, cifra oficial, informe, albergue, rescate, fosa clandestina, agresión, búsqueda en campo...).",
    "fecha: formato AAAA-MM-DD. Si solo se conoce el mes, usa el día 01 y anótalo en descripcion.",
    "lat, lon: grados decimales. En Google Maps clic derecho sobre el punto y copiar coordenadas. Si solo se conoce el municipio, usa la cabecera y anota 'ubicación aproximada' en descripcion.",
    "lugar: localidad, municipio y estado, por ejemplo 'Tepuche, Culiacán, Sinaloa'. El nombre del estado se usa para contar eventos por entidad.",
    "titulo: una línea. descripcion: qué pasó según la fuente, con cifras y quién las dio, sin datos que identifiquen a personas.",
    "fuente: id del catálogo (hoja 'catalogos'). Para agregar una fuente nueva hay que darla de alta en datos/fuentes.json y regenerar esta plantilla.",
    "url: enlace a la nota o informe, si existe.",
    "verificacion: oficial (la cifra la dio una autoridad), organización (la documentó una organización civil), campo, prensa (solo la nota) o sin verificar.",
    "ejemplo: si / no.",
    "",
    "Para publicar: guarda este archivo, cópialo a la raíz del repositorio y ejecuta",
    "    python herramientas/actualizar_eventos.py captura_eventos.xlsx",
    "El script valida las filas, actualiza datos/eventos.json y deja un respaldo del JSON anterior. Luego sube datos/eventos.json a GitHub.",
    "Para regenerar esta plantilla con los eventos y fuentes actuales: python herramientas/crear_plantilla.py",
]


def main():
    salida = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("captura_eventos.xlsx")
    eventos = json.loads(Path("datos/eventos.json").read_text(encoding="utf-8"))
    fuentes = json.loads(Path("datos/fuentes.json").read_text(encoding="utf-8"))
    arial, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    amarillo, azul = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9E2F3")

    wb = Workbook()
    we = wb.active
    we.title = "eventos"
    we.append(COLS)
    for e in sorted(eventos, key=lambda x: (x["fecha"], x["id"])):
        we.append([("si" if e[c] else "no") if c == "ejemplo" else e.get(c, "") for c in COLS])
    for i, w in enumerate(ANCHOS, 1):
        we.column_dimensions[get_column_letter(i)].width = w
    for row in we.iter_rows():
        for c in row:
            c.font = arial
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (7, 8, 9))
    for c in we[1]:
        c.font, c.fill = bold, azul
    for r in range(2, MAX_FILAS + 2):
        for ci in range(1, len(COLS) + 1):
            we.cell(r, ci).fill = amarillo
        we.cell(r, 4).number_format = "yyyy-mm-dd"
    we.freeze_panes = "B2"

    wc = wb.create_sheet("catalogos")
    wc.append(["tema", "verificacion", "fuente_id", "fuente_nombre", "fuente_tipo"])
    n = max(len(TEMAS), len(VERIF), len(fuentes))
    for i in range(n):
        f = fuentes[i] if i < len(fuentes) else None
        wc.append([TEMAS[i] if i < len(TEMAS) else None, VERIF[i] if i < len(VERIF) else None,
                   f["id"] if f else None, f["nombre"] if f else None, f["tipo"] if f else None])
    for w, col in zip([16, 16, 18, 80, 14], "ABCDE"):
        wc.column_dimensions[col].width = w
    for row in wc.iter_rows():
        for c in row:
            c.font = arial
    for c in wc[1]:
        c.font, c.fill = bold, azul

    rango = f"2:{MAX_FILAS + 1}"
    reglas = [
        DataValidation(type="list", formula1=f"=catalogos!$A$2:$A${len(TEMAS) + 1}", allow_blank=True), "B",
        DataValidation(type="list", formula1=f"=catalogos!$B$2:$B${len(VERIF) + 1}", allow_blank=True), "L",
        DataValidation(type="list", formula1=f"=catalogos!$C$2:$C${len(fuentes) + 1}", allow_blank=True), "J",
        DataValidation(type="list", formula1='"si,no"', allow_blank=True), "M",
        DataValidation(type="decimal", operator="between", formula1="14", formula2="33", allow_blank=True), "E",
        DataValidation(type="decimal", operator="between", formula1="-119", formula2="-86", allow_blank=True), "F",
    ]
    for dv, col in zip(reglas[::2], reglas[1::2]):
        dv.error, dv.showErrorMessage = "Valor no permitido", True
        dv.add(f"{col}{rango.replace(':', ':' + col)}")
        we.add_data_validation(dv)

    wi = wb.create_sheet("instrucciones")
    for t in INSTRUCCIONES:
        wi.append([t])
    for row in wi.iter_rows():
        for c in row:
            c.font = bold if c.row == 1 else arial
    wi.column_dimensions["A"].width = 130
    wb.save(salida)
    print(f"Plantilla escrita en {salida} con {len(eventos)} eventos y {len(fuentes)} fuentes")


if __name__ == "__main__":
    main()
